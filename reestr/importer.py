from __future__ import annotations

import hashlib
import io
import re
from datetime import datetime
from pathlib import Path
from typing import Any

import xlrd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .domain import (
    DATE_FIELDS,
    MONEY_FIELDS,
    TEXT_FIELDS,
    ValidationError,
    minor_to_number,
    normalize_contract,
    parse_iso_date,
)


MAX_IMPORT_ROWS = 10_000


class ImportFileError(ValueError):
    pass


def normalize_header(value: Any) -> str:
    text = str(value or "").strip().casefold().replace("ё", "е")
    return re.sub(r"[^a-zа-я0-9]+", " ", text).strip()


HEADER_ALIASES = {
    "party_type": (
        "к клиент п подрядчик",
        "к клиент п подрядчик",
        "тип контрагента",
        "клиент подрядчик",
        "к п",
    ),
    "counterparty": ("наименование контрагента", "контрагент", "наименование"),
    "country": ("страна",),
    "contract": ("и дата договора", "договор", "номер и дата договора", "договор и дата"),
    "request": ("и дата заявки", "заявки и дата", "заявка", "номер и дата заявки", "заявка и дата"),
    "invoice": ("и дата счета", "счета и дата", "счет", "номер и дата счета", "счет и дата"),
    "invoice_status": (
        "статус счета о оплачен н не оплачен",
        "статус счета",
        "оплачен",
        "статус оплаты",
    ),
    "payment_date": ("дата оплаты", "фактическая дата оплаты"),
    "payment_due_date": ("сроки оплаты", "срок оплаты", "дата платежа", "плановая дата оплаты"),
    "avr_number": ("авр", "номер авр", "акт", "номер акта"),
    "avr_date": ("дата авр", "дата акта"),
    "avr_amount": ("сумма авр", "сумма акта", "сумма"),
    "vat_amount": (
        "в том числе сумма ндс авр",
        "сумма ндс авр",
        "в том числе ндс",
        "сумма ндс",
        "ндс",
    ),
    "currency": ("валюта авр", "валюта", "валюта акта"),
    "client": ("под клиента", "клиент", "для клиента"),
    "notes": ("примечания", "примечание", "комментарий"),
    "payment_reference": (
        "номер п п сумма оплаты",
        "номер пп сумма оплаты",
        "номер п п",
        "платежное поручение",
        "платежка",
    ),
}

HEADER_TO_FIELD = {
    normalize_header(alias): field for field, aliases in HEADER_ALIASES.items() for alias in aliases
}

SAMPLE_HEADERS = [
    "К- клиент, П -подрядчик",
    "Наименование контрагента",
    "Страна",
    "№ и дата договора",
    "№ заявки и дата",
    "№ счета и дата",
    "Статус счета: О-оплачен, Н-не оплачен",
    "дата оплаты",
    "СРОКИ ОПЛАТЫ",
    "№ АВР",
    "Дата АВР",
    "Сумма АВР",
    "в том числе сумма НДС АВР",
    "Валюта АВР",
    "Под КЛИЕНТА",
    "примечания",
    "Номер п/п, сумма оплаты",
]


def _xlsx_rows(content: bytes) -> tuple[str, list[list[Any]]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        return sheet.title, rows
    except Exception as exc:
        raise ImportFileError("Не удалось прочитать XLSX. Файл повреждён или имеет другой формат.") from exc


def _xls_rows(content: bytes) -> tuple[str, list[list[Any]]]:
    try:
        workbook = xlrd.open_workbook(file_contents=content)
        sheet = workbook.sheet_by_index(0)
        rows: list[list[Any]] = []
        for row_index in range(sheet.nrows):
            row = []
            for col_index in range(sheet.ncols):
                cell = sheet.cell(row_index, col_index)
                value = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate.xldate_as_datetime(value, workbook.datemode)
                elif cell.ctype == xlrd.XL_CELL_EMPTY:
                    value = None
                row.append(value)
            rows.append(row)
        return sheet.name, rows
    except Exception as exc:
        raise ImportFileError("Не удалось прочитать XLS. Файл повреждён или имеет другой формат.") from exc


def _find_headers(rows: list[list[Any]]) -> tuple[int, dict[int, str], list[str]]:
    best: tuple[int, dict[int, str], list[str]] | None = None
    best_count = 0
    for row_index, row in enumerate(rows[:30]):
        titles = [str(value or "").strip() for value in row]
        mapping = {
            index: HEADER_TO_FIELD[normalize_header(value)]
            for index, value in enumerate(row)
            if normalize_header(value) in HEADER_TO_FIELD
        }
        if len(mapping) > best_count:
            best = (row_index, mapping, titles)
            best_count = len(mapping)
    if best is None or best_count < 2:
        raise ImportFileError(
            "Не найдена строка заголовков. Используйте файл по образцу «Реестр договоров1.xlsx»."
        )
    required = {"counterparty", "contract"}
    if not required.issubset(best[1].values()):
        raise ImportFileError("В файле должны быть колонки «Наименование контрагента» и «№ и дата договора».")
    return best


def _split_number_date(value: Any, prefix: str, errors: dict[str, str]) -> tuple[str, str | None]:
    if value is None or value == "":
        return "", None
    if isinstance(value, datetime):
        return "", value.date().isoformat()
    text = str(value).strip()
    match = re.match(
        r"^(.*?)\s+от\s+(\d{1,2}[./-]\d{1,2}[./-]\d{2,4})\s*$",
        text,
        flags=re.IGNORECASE,
    )
    if not match:
        return text, None
    number = match.group(1).strip()
    try:
        parsed = parse_iso_date(match.group(2))
    except ValueError:
        errors[f"{prefix}_date"] = "Не удалось распознать дату после слова «от»"
        return text, None
    return number, parsed


def _public_data(data: dict[str, Any]) -> dict[str, Any]:
    result = dict(data)
    result["avr_amount"] = minor_to_number(result.pop("avr_amount_minor", 0))
    result["vat_amount"] = minor_to_number(result.pop("vat_amount_minor", 0))
    return result


def parse_import(filename: str, content: bytes) -> dict[str, Any]:
    extension = Path(filename).suffix.casefold()
    if extension not in {".xlsx", ".xls"}:
        raise ImportFileError("Поддерживаются только файлы .xlsx и .xls")
    if not content:
        raise ImportFileError("Файл пуст")
    if len(content) > 15 * 1024 * 1024:
        raise ImportFileError("Размер файла не должен превышать 15 МБ")

    sheet_name, rows = _xlsx_rows(content) if extension == ".xlsx" else _xls_rows(content)
    header_index, mapping, header_titles = _find_headers(rows)
    result_rows = []
    for row_index in range(header_index + 1, len(rows)):
        if len(result_rows) >= MAX_IMPORT_ROWS:
            raise ImportFileError(f"За один раз можно импортировать не более {MAX_IMPORT_ROWS} строк")
        source = rows[row_index]
        if not any(value is not None and str(value).strip() for value in source):
            continue
        raw: dict[str, Any] = {
            "party_type": "К",
            "currency": "RUB",
            "invoice_status": "Н",
        }
        errors: dict[str, str] = {}
        for column_index, field in mapping.items():
            value = source[column_index] if column_index < len(source) else None
            if field in {"contract", "request", "invoice"}:
                prefix = {"contract": "contract", "request": "request", "invoice": "invoice"}[field]
                number, parsed_date = _split_number_date(value, prefix, errors)
                raw[f"{prefix}_number"] = number
                raw[f"{prefix}_date"] = parsed_date
            elif field == "payment_due_date":
                try:
                    raw["payment_due_date"] = parse_iso_date(value)
                    raw["payment_terms"] = "" if raw["payment_due_date"] else ""
                except ValueError:
                    # The source column often contains text such as "30 дней".
                    raw["payment_due_date"] = None
                    raw["payment_terms"] = str(value).strip()
            else:
                raw[field] = value
        try:
            normalized = normalize_contract(raw)
        except ValidationError as exc:
            errors.update(exc.errors)
            # Keep a safe, readable preview even when the source row is invalid.
            safe = dict(raw)
            safe["counterparty"] = str(raw.get("counterparty") or "—")
            safe["contract_number"] = str(raw.get("contract_number") or "—")
            for field in errors:
                if field in DATE_FIELDS:
                    safe[field] = None
                elif field in MONEY_FIELDS:
                    safe[field] = 0
                elif field == "party_type":
                    safe[field] = "К"
                elif field == "invoice_status":
                    safe[field] = "Н"
                elif field == "currency":
                    safe[field] = "RUB"
                elif field in TEXT_FIELDS:
                    safe[field] = str(raw.get(field) or "")[: TEXT_FIELDS[field]]
                    if field in {"counterparty", "contract_number"} and not safe[field]:
                        safe[field] = "—"
            normalized = normalize_contract(safe)
        item = {
            "source_row": row_index + 1,
            "data": _public_data(normalized),
            "errors": errors,
        }
        result_rows.append(item)

    if not result_rows:
        raise ImportFileError("После строки заголовков нет данных для импорта")
    valid = sum(not item["errors"] for item in result_rows)
    return {
        "filename": Path(filename).name,
        "sha256": hashlib.sha256(content).hexdigest(),
        "sheet": sheet_name,
        "header_row": header_index + 1,
        "headers": [
            {
                "column": index + 1,
                "title": title or f"Колонка {index + 1}",
                "field": mapping.get(index),
            }
            for index, title in enumerate(header_titles)
        ],
        "rows": result_rows,
        "total": len(result_rows),
        "valid": valid,
        "invalid": len(result_rows) - valid,
    }


def make_import_template() -> io.BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Реестр договоров"
    sheet.append(SAMPLE_HEADERS)
    sheet.append(
        [
            "К",
            "Пример контрагента",
            "РФ",
            "К-101 от 01.03.2026",
            "15 от 03.03.2026",
            "К-101-1 от 05.03.2026",
            "Н",
            None,
            "31.03.2026",
            "К-101-1",
            "05.03.2026",
            12000.00,
            2000.00,
            "RUB",
            "Основной клиент",
            "Строку-пример перед импортом удалите",
            "",
        ]
    )
    fill = PatternFill("solid", fgColor="1F5D73")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 45
    widths = [20, 32, 12, 25, 24, 27, 28, 16, 18, 18, 15, 16, 24, 14, 25, 35, 32]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:Q{sheet.max_row}"
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
