from __future__ import annotations

import html
import io
from datetime import datetime
from decimal import Decimal
from typing import Any

import xlwt
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .domain import display_date
from .resources import resource_path


FILTER_NAMES = {
    "counterparty": "Контрагент",
    "currency": "Валюта",
    "date_from": "Период с",
    "date_to": "Период по",
    "status": "Статус",
}


def format_value(value: Any, kind: str = "text") -> str:
    if value is None or value == "":
        return "—"
    if kind == "date":
        return display_date(str(value))
    if kind == "money":
        amount = Decimal(str(value)).quantize(Decimal("0.01"))
        return f"{amount:,.2f}".replace(",", " ").replace(".", ",")
    if kind == "integer":
        return str(int(value))
    return str(value)


def generated_text(value: str) -> str:
    try:
        return datetime.fromisoformat(value).strftime("%d.%m.%Y %H:%M")
    except ValueError:
        return value


def report_metadata(report: dict[str, Any]) -> list[str]:
    lines = [f"Сформирован: {generated_text(report['generated_at'])}", f"Строк: {report['row_count']}"]
    for key, value in report.get("filters", {}).items():
        label = FILTER_NAMES.get(key, key)
        if key in {"date_from", "date_to"}:
            value = display_date(value)
        lines.append(f"{label}: {value}")
    return lines


def export_xls(report: dict[str, Any]) -> io.BytesIO:
    workbook = xlwt.Workbook(encoding="utf-8")
    sheet = workbook.add_sheet("Отчёт")
    columns = report["columns"]
    title_style = xlwt.easyxf("font: name Arial, bold on, height 320; align: vert center")
    meta_style = xlwt.easyxf("font: name Arial, height 180, color gray50")
    header_style = xlwt.easyxf(
        "font: name Arial, bold on, color white; pattern: pattern solid, fore_colour 0x35; "
        "align: wrap on, horiz center, vert center; borders: bottom thin, top thin, left thin, right thin"
    )
    text_style = xlwt.easyxf(
        "font: name Arial, height 180; align: wrap on, vert top; borders: bottom hair, left hair, right hair"
    )
    money_style = xlwt.easyxf(
        "font: name Arial, height 180; align: horiz right, vert top; borders: bottom hair, left hair, right hair",
        num_format_str="#,##0.00",
    )
    integer_style = xlwt.easyxf(
        "font: name Arial, height 180; align: horiz right; borders: bottom hair, left hair, right hair",
        num_format_str="0",
    )

    last_column = max(0, len(columns) - 1)
    sheet.write_merge(0, 0, 0, last_column, report["title"], title_style)
    sheet.row(0).height_mismatch = True
    sheet.row(0).height = 480
    row_index = 1
    for line in report_metadata(report):
        sheet.write_merge(row_index, row_index, 0, last_column, line, meta_style)
        row_index += 1
    row_index += 1
    header_row = row_index
    for column_index, column in enumerate(columns):
        sheet.write(header_row, column_index, column["title"], header_style)
        sheet.col(column_index).width = min(65_535, int(column.get("width", 16)) * 270)
    sheet.row(header_row).height_mismatch = True
    sheet.row(header_row).height = 620

    for item in report["rows"]:
        row_index += 1
        for column_index, column in enumerate(columns):
            value = item.get(column["key"])
            kind = column.get("kind", "text")
            if kind == "money":
                sheet.write(row_index, column_index, float(value or 0), money_style)
            elif kind == "integer":
                sheet.write(row_index, column_index, int(value or 0), integer_style)
            else:
                sheet.write(row_index, column_index, format_value(value, kind), text_style)

    sheet.set_panes_frozen(True)
    sheet.set_horz_split_pos(header_row + 1)
    sheet.set_remove_splits(True)
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _shade_word_cell(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _repeat_word_header(row: Any) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def export_docx(report: dict[str, Any]) -> io.BytesIO:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.0)
    section.right_margin = Cm(1.0)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(8)

    title = document.add_paragraph()
    title_run = title.add_run(report["title"])
    title_run.bold = True
    title_run.font.size = Pt(16)
    title_run.font.color.rgb = RGBColor(31, 93, 115)
    for line in report_metadata(report):
        paragraph = document.add_paragraph(line)
        paragraph.paragraph_format.space_after = Pt(1)
        for run in paragraph.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(90, 100, 105)

    columns = report["columns"]
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    table.autofit = False
    _repeat_word_header(table.rows[0])
    total_width = sum(column.get("width", 15) for column in columns)
    available_width = 27.7
    widths = [Cm(available_width * column.get("width", 15) / total_width) for column in columns]

    for index, column in enumerate(columns):
        cell = table.rows[0].cells[index]
        cell.width = widths[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _shade_word_cell(cell, "1F5D73")
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(column["title"])
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(7)

    for item in report["rows"]:
        cells = table.add_row().cells
        for index, column in enumerate(columns):
            cell = cells[index]
            cell.width = widths[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            paragraph = cell.paragraphs[0]
            kind = column.get("kind", "text")
            if kind in {"money", "integer", "date"}:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = paragraph.add_run(format_value(item.get(column["key"]), kind))
            run.font.size = Pt(7)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Реестр договоров  •  ")
    run.font.size = Pt(8)
    page_field = OxmlElement("w:fldSimple")
    page_field.set(qn("w:instr"), "PAGE")
    footer._p.append(page_field)
    document.core_properties.title = report["title"]
    document.core_properties.subject = "Реестр договоров"
    output = io.BytesIO()
    document.save(output)
    output.seek(0)
    return output


_FONT_REGISTERED = False


def register_pdf_fonts() -> None:
    global _FONT_REGISTERED
    if _FONT_REGISTERED:
        return
    regular = resource_path("reestr/assets/fonts/NotoSans-Regular.ttf")
    bold = resource_path("reestr/assets/fonts/NotoSans-Bold.ttf")
    pdfmetrics.registerFont(TTFont("ReestrSans", str(regular)))
    pdfmetrics.registerFont(TTFont("ReestrSans-Bold", str(bold)))
    pdfmetrics.registerFontFamily(
        "ReestrSans",
        normal="ReestrSans",
        bold="ReestrSans-Bold",
        italic="ReestrSans",
        boldItalic="ReestrSans-Bold",
    )
    _FONT_REGISTERED = True


def export_pdf(report: dict[str, Any]) -> io.BytesIO:
    register_pdf_fonts()
    output = io.BytesIO()
    page_size = landscape(A4)
    document = SimpleDocTemplate(
        output,
        pagesize=page_size,
        leftMargin=9 * mm,
        rightMargin=9 * mm,
        topMargin=9 * mm,
        bottomMargin=12 * mm,
        title=report["title"],
        author="Реестр договоров",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Title"],
        fontName="ReestrSans-Bold",
        fontSize=15,
        leading=18,
        textColor=colors.HexColor("#1F5D73"),
        spaceAfter=5,
    )
    meta_style = ParagraphStyle(
        "ReportMeta",
        parent=styles["Normal"],
        fontName="ReestrSans",
        fontSize=7.5,
        leading=10,
        textColor=colors.HexColor("#5A6469"),
    )
    cell_style = ParagraphStyle(
        "ReportCell",
        parent=styles["Normal"],
        fontName="ReestrSans",
        fontSize=6.2,
        leading=7.6,
        splitLongWords=True,
    )
    number_style = ParagraphStyle("ReportNumber", parent=cell_style, alignment=TA_RIGHT)
    header_style = ParagraphStyle(
        "ReportHeader",
        parent=cell_style,
        fontName="ReestrSans-Bold",
        textColor=colors.white,
        alignment=1,
    )
    story = [Paragraph(html.escape(report["title"]), title_style)]
    story.extend(Paragraph(html.escape(line), meta_style) for line in report_metadata(report))
    story.append(Spacer(1, 4 * mm))

    columns = report["columns"]
    data = [
        [Paragraph(html.escape(column["title"]), header_style) for column in columns]
    ]
    for item in report["rows"]:
        row = []
        for column in columns:
            kind = column.get("kind", "text")
            style = number_style if kind in {"money", "integer", "date"} else cell_style
            value = format_value(item.get(column["key"]), kind)
            row.append(Paragraph(html.escape(value), style))
        data.append(row)

    total_width = sum(column.get("width", 15) for column in columns)
    available_width = page_size[0] - 18 * mm
    widths = [available_width * column.get("width", 15) / total_width for column in columns]
    table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F5D73")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("VALIGN", (0, 0), (-1, 0), "MIDDLE"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C7D0D3")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F7F8")]),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    story.append(table)

    def footer(canvas: Any, doc: Any) -> None:
        canvas.saveState()
        canvas.setFont("ReestrSans", 7)
        canvas.setFillColor(colors.HexColor("#6B7478"))
        canvas.drawString(9 * mm, 6 * mm, "Реестр договоров")
        canvas.drawRightString(page_size[0] - 9 * mm, 6 * mm, f"Страница {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    output.seek(0)
    return output


def export_report(report: dict[str, Any], file_format: str) -> tuple[io.BytesIO, str, str]:
    if file_format == "xls":
        return export_xls(report), "application/vnd.ms-excel", "xls"
    if file_format == "docx":
        return (
            export_docx(report),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "docx",
        )
    if file_format == "pdf":
        return export_pdf(report), "application/pdf", "pdf"
    raise ValueError("Поддерживаются форматы xls, docx и pdf")
