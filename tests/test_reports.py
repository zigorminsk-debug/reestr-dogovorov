from __future__ import annotations

import sqlite3
from datetime import date, timedelta
from io import BytesIO

import xlrd
from docx import Document


def seed(client, headers, payload):
    today = date.today()
    rows = [
        ("Альфа", "Н", (today - timedelta(days=4)).isoformat(), None, "1000.10"),
        ("Альфа", "О", (today - timedelta(days=15)).isoformat(), (today - timedelta(days=10)).isoformat(), "2000.20"),
        ("Бета", "О", (today + timedelta(days=5)).isoformat(), today.isoformat(), "3000.30"),
    ]
    for index, (counterparty, status, due, paid, amount) in enumerate(rows):
        item = dict(
            payload,
            counterparty=counterparty,
            contract_number=f"Д-{index + 1}",
            invoice_status=status,
            payment_due_date=due,
            payment_date=paid,
            avr_amount=amount,
            vat_amount="0",
        )
        assert client.post("/api/contracts", json=item, headers=headers).status_code == 201


def test_report_calculations_and_filters(client, headers, contract_payload):
    seed(client, headers, contract_payload)
    counterparties = client.get("/api/reports/counterparties").json
    assert counterparties["row_count"] == 2
    alpha = next(row for row in counterparties["rows"] if row["counterparty"] == "Альфа")
    assert alpha["contract_count"] == 2
    assert alpha["avr_amount"] == 3000.30
    assert alpha["paid_amount"] == 2000.20
    assert alpha["outstanding_amount"] == 1000.10
    assert alpha["overdue_count"] == 1
    assert alpha["paid_late_count"] == 1

    overdue = client.get("/api/reports/overdue").json
    assert overdue["row_count"] == 2
    unpaid = client.get("/api/reports/overdue?status=overdue_unpaid").json
    assert unpaid["row_count"] == 1
    assert unpaid["rows"][0]["counterparty"] == "Альфа"
    payments = client.get("/api/reports/payments?status=paid&counterparty=Бета").json
    assert payments["row_count"] == 1
    assert payments["rows"][0]["avr_amount"] == 3000.30
    assert client.get("/api/reports/payments?date_from=not-a-date").status_code == 422


def test_all_export_formats_are_real_files(client, headers, contract_payload):
    seed(client, headers, contract_payload)
    xls = client.get("/api/reports/counterparties/export/xls")
    assert xls.status_code == 200
    assert xls.data[:8] == bytes.fromhex("D0CF11E0A1B11AE1")
    workbook = xlrd.open_workbook(file_contents=xls.data)
    sheet = workbook.sheet_by_index(0)
    assert sheet.cell_value(0, 0) == "Отчёт по контрагентам"
    assert any(sheet.cell_value(row, 0) == "Альфа" for row in range(sheet.nrows))

    docx = client.get("/api/reports/payments/export/docx")
    assert docx.status_code == 200
    document = Document(BytesIO(docx.data))
    assert document.paragraphs[0].text == "Отчёт по срокам оплаты"
    assert len(document.tables) == 1
    assert len(document.tables[0].rows) == 4

    pdf = client.get("/api/reports/overdue/export/pdf")
    assert pdf.status_code == 200
    assert pdf.data.startswith(b"%PDF-")
    assert len(pdf.data) > 10_000
    assert b"NotoSans" in pdf.data
    assert b"/FontFile2" in pdf.data  # Cyrillic TTF is embedded, not substituted.


def test_sqlite_backup_contains_shared_data(client, headers, contract_payload, tmp_path):
    assert client.post("/api/contracts", json=contract_payload, headers=headers).status_code == 201
    response = client.get("/api/backup")
    assert response.status_code == 200
    target = tmp_path / "downloaded-backup.db"
    target.write_bytes(response.data)
    connection = sqlite3.connect(target)
    try:
        assert connection.execute("SELECT count(*) FROM contracts").fetchone()[0] == 1
    finally:
        connection.close()
